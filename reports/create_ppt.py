"""生成天工系统论文创新点汇报PPT"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ── 颜色定义 ──
C_DARK    = RGBColor(0x2C, 0x3E, 0x50)
C_BLUE    = RGBColor(0x17, 0x4E, 0xA6)
C_GREEN   = RGBColor(0x0D, 0x65, 0x2D)
C_PURPLE  = RGBColor(0x68, 0x1D, 0xA8)
C_ORANGE  = RGBColor(0xB0, 0x60, 0x00)
C_RED     = RGBColor(0xC5, 0x22, 0x1F)
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_LGRAY   = RGBColor(0xF0, 0xF0, 0xF0)
C_GRAY    = RGBColor(0x5F, 0x63, 0x68)
C_ACCENT  = RGBColor(0x1A, 0x73, 0xE8)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW = prs.slide_width
SH = prs.slide_height


def add_blank_slide():
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def add_bg(slide, color=C_WHITE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, fill_color=None, line_color=None, line_width=Pt(1)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color or C_WHITE
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = line_width
    else:
        shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text, font_size=18,
                bold=False, color=C_DARK, align=PP_ALIGN.LEFT, font_name="Microsoft YaHei"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = align
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=16,
                    color=C_DARK, spacing=Pt(8), bold_first=False):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Microsoft YaHei"
        p.space_after = spacing
        p.level = 0
        if bold_first and i == 0:
            p.font.bold = True
    return txBox


def add_card(slide, left, top, width, height, title, bullets, title_color, bg_color):
    """添加一个卡片（标题+项目列表）"""
    card = add_rect(slide, left, top, width, height, fill_color=bg_color, line_color=title_color, line_width=Pt(2))
    add_textbox(slide, left + Inches(0.2), top + Inches(0.1), width - Inches(0.4), Inches(0.5),
                title, font_size=20, bold=True, color=title_color)
    bullet_top = top + Inches(0.55)
    add_bullet_list(slide, left + Inches(0.25), bullet_top, width - Inches(0.5),
                    height - Inches(0.65), bullets, font_size=14, color=C_DARK, spacing=Pt(4))
    return card


# ════════════════════════════════════════════
# Slide 1: 封面
# ════════════════════════════════════════════
sl = add_blank_slide()
add_rect(sl, Inches(0), Inches(0), SW, SH, fill_color=RGBColor(0x1A, 0x73, 0xE8))
add_textbox(sl, Inches(1), Inches(1.8), Inches(11.3), Inches(1.2),
            "天工系统（SkyEngine）", font_size=44, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
add_textbox(sl, Inches(1), Inches(3.0), Inches(11.3), Inches(0.8),
            "面向 FJSP-MAPF 联合调度的强化学习环境\n与多专家蒸馏学习框架",
            font_size=26, color=RGBColor(0xBB, 0xDE, 0xFB), align=PP_ALIGN.CENTER)
add_textbox(sl, Inches(1), Inches(5.5), Inches(11.3), Inches(0.5),
            "论文创新点汇报", font_size=20, color=RGBColor(0x90, 0xCA, 0xF9), align=PP_ALIGN.CENTER)

# ════════════════════════════════════════════
# Slide 2: 核心问题 — 为什么要做这件事
# ════════════════════════════════════════════
sl = add_blank_slide()
add_rect(sl, Inches(0), Inches(0), SW, Inches(0.9), fill_color=C_DARK)
add_textbox(sl, Inches(0.5), Inches(0.15), Inches(12), Inches(0.6),
            "核心问题：FJSP 与 MAPF 的耦合损失", font_size=28, bold=True, color=C_WHITE, align=PP_ALIGN.LEFT)

# 左：两大模块
add_card(sl, Inches(0.5), Inches(1.2), Inches(5.8), Inches(2.0),
         "两大独立模块，信息隔离",
         ["FJSP（柔性调度）：决策加工机器与工序顺序，仅优化加工时间",
          "MAPF（路径规划）：决策 AGV 搬运路径，仅优化通行效率",
          "搬运衔接环节无人统筹 → 严重耦合损失"],
         C_BLUE, RGBColor(0xE8, 0xF0, 0xFE))

# 右：三类损失
add_card(sl, Inches(6.8), Inches(1.2), Inches(6.0), Inches(2.0),
         "三类核心耦合损失",
         ["❶ 距离盲区：选最快机器但 AGV 远距离搬运，总工期反增",
          "❷ 时序错位：送达时间与运输时序不匹配，机器空闲或物料堆积",
          "❸ 拥堵聚集：多 AGV 抢占窄通道，引发阻塞卡顿"],
         C_RED, RGBColor(0xFC, 0xE8, 0xE6))

# 下方：现有方案局限
add_card(sl, Inches(0.5), Inches(3.6), Inches(12.3), Inches(1.8),
         "现有方案的共同局限",
         ["MAPF 被简化为运输时间矩阵，不模拟实际路径冲突、阻塞、绕行",
          "耦合被视为单向约束（FJSP→MAPF），不建模路径结果对调度的反向影响",
          "不存在统一环境将「调度产生路径需求 → 路径结果重塑调度机会」的闭环显式展开"],
         C_GRAY, C_LGRAY)

# 底部：论文故事线
add_rect(sl, Inches(0.5), Inches(5.8), Inches(12.3), Inches(1.2),
         fill_color=RGBColor(0xE8, 0xF0, 0xFE), line_color=C_BLUE, line_width=Pt(2))
add_textbox(sl, Inches(0.7), Inches(5.9), Inches(12), Inches(0.4),
            "论文故事线", font_size=18, bold=True, color=C_BLUE)
add_textbox(sl, Inches(0.7), Inches(6.3), Inches(12), Inches(0.5),
            "现象（三类耦合损失）→ 分析（Assigner 是核心接口）→ 环境（SkyEngine 统一闭环）→ 方法（多专家蒸馏 + 多塔网络）→ 验证",
            font_size=15, color=C_DARK)

# ════════════════════════════════════════════
# Slide 3: 三个创新点总览
# ════════════════════════════════════════════
sl = add_blank_slide()
add_rect(sl, Inches(0), Inches(0), SW, Inches(0.9), fill_color=C_DARK)
add_textbox(sl, Inches(0.5), Inches(0.15), Inches(12), Inches(0.6),
            "三条创新点：环境 → 框架 → 网络，逐层递进", font_size=28, bold=True, color=C_WHITE)

# 三个卡片并排
card_w = Inches(3.8)
card_h = Inches(5.0)
gap = Inches(0.45)
start_x = Inches(0.7)
card_y = Inches(1.3)

# 创新点一
add_rect(sl, start_x, card_y, card_w, card_h,
         fill_color=RGBColor(0xE8, 0xF0, 0xFE), line_color=C_BLUE, line_width=Pt(2.5))
add_textbox(sl, start_x + Inches(0.2), card_y + Inches(0.15), card_w - Inches(0.4), Inches(0.4),
            "创新点一", font_size=16, bold=True, color=C_BLUE)
add_textbox(sl, start_x + Inches(0.2), card_y + Inches(0.5), card_w - Inches(0.4), Inches(0.6),
            "SkyEngine\n闭环统一 RL 环境", font_size=22, bold=True, color=C_DARK)
add_bullet_list(sl, start_x + Inches(0.25), card_y + Inches(1.2), card_w - Inches(0.5), card_h - Inches(1.4),
                ["闭环共仿真：FJSP 调度与 MAPF 路由在共享时间线上增量推进",
                 "显式阻塞反馈：AGV 阻塞、绕行作为实际事件反向传播",
                 "可插拔 Assigner 模块：调度与路径之间的核心桥梁",
                 "17 项指标的评估体系"],
                font_size=13, color=C_DARK, spacing=Pt(10))

# 箭头 1→2
arrow_x1 = start_x + card_w + Inches(0.05)
add_textbox(sl, arrow_x1, card_y + card_h / 2 - Inches(0.2), Inches(0.4), Inches(0.4),
            "→", font_size=28, bold=True, color=C_GRAY, align=PP_ALIGN.CENTER)

# 创新点二
x2 = start_x + card_w + gap
add_rect(sl, x2, card_y, card_w, card_h,
         fill_color=RGBColor(0xE6, 0xF4, 0xEA), line_color=C_GREEN, line_width=Pt(2.5))
add_textbox(sl, x2 + Inches(0.2), card_y + Inches(0.15), card_w - Inches(0.4), Inches(0.4),
            "创新点二", font_size=16, bold=True, color=C_GREEN)
add_textbox(sl, x2 + Inches(0.2), card_y + Inches(0.5), card_w - Inches(0.4), Inches(0.6),
            "多专家在线蒸馏\nStudent-Induced 框架", font_size=22, bold=True, color=C_DARK)
add_bullet_list(sl, x2 + Inches(0.25), card_y + Inches(1.2), card_w - Inches(0.5), card_h - Inches(1.4),
                ["在 student 自己诱导出的状态上蒸馏，消除分布偏移",
                 "多专家可靠性动态评估，状态依赖权重",
                 "在线贝叶斯更新专家可靠性",
                 "蒸馏稳定性保障：小头部网络确保微调稳定"],
                font_size=13, color=C_DARK, spacing=Pt(10))

# 箭头 2→3
arrow_x2 = x2 + card_w + Inches(0.05)
add_textbox(sl, arrow_x2, card_y + card_h / 2 - Inches(0.2), Inches(0.4), Inches(0.4),
            "→", font_size=28, bold=True, color=C_GRAY, align=PP_ALIGN.CENTER)

# 创新点三
x3 = x2 + card_w + gap
add_rect(sl, x3, card_y, card_w, card_h,
         fill_color=RGBColor(0xF3, 0xE8, 0xFD), line_color=C_PURPLE, line_width=Pt(2.5))
add_textbox(sl, x3 + Inches(0.2), card_y + Inches(0.15), card_w - Inches(0.4), Inches(0.4),
            "创新点三", font_size=16, bold=True, color=C_PURPLE)
add_textbox(sl, x3 + Inches(0.2), card_y + Inches(0.5), card_w - Inches(0.4), Inches(0.6),
            "专家增强型\n多塔策略网络", font_size=22, bold=True, color=C_DARK)
add_bullet_list(sl, x3 + Inches(0.25), card_y + Inches(1.2), card_w - Inches(0.5), card_h - Inches(1.4),
                ["状态塔：AGV、机器、Job、搬运任务多维特征",
                 "专家塔：冻结预训练参数 + 轻量 Adapter",
                 "动态路由融合：w_k(s) 指导专家特征加权",
                 "从专家评估到专家融合的完整闭环"],
                font_size=13, color=C_DARK, spacing=Pt(10))

# ════════════════════════════════════════════
# Slide 4: 创新点一详解 — SkyEngine
# ════════════════════════════════════════════
sl = add_blank_slide()
add_rect(sl, Inches(0), Inches(0), SW, Inches(0.9), fill_color=C_BLUE)
add_textbox(sl, Inches(0.5), Inches(0.15), Inches(12), Inches(0.6),
            "创新点一：SkyEngine — 首个面向 FJSP-MAPF 联合调度的闭环 RL 环境",
            font_size=24, bold=True, color=C_WHITE)

# 三大创新
add_card(sl, Inches(0.5), Inches(1.2), Inches(3.9), Inches(3.5),
         "① 闭环共仿真",
         ["整合 FJSP 调度与 MAPF 路由，共享时间线同步推进",
          "非「先调度完再执行路径」，而是增量决策",
          "每个调度动作立即与当前 AGV 和机器状态交互"],
         C_BLUE, RGBColor(0xE8, 0xF0, 0xFE))

add_card(sl, Inches(4.8), Inches(1.2), Inches(3.9), Inches(3.5),
         "② 显式阻塞反馈与状态建模",
         ["AGV 阻塞、等待、绕行不被抽象为时间惩罚",
          "作为实际事件抽取特征，反向传播影响调度选择",
          "统一观测接口 + 17 项指标评估体系"],
         C_BLUE, RGBColor(0xE8, 0xF0, 0xFE))

add_card(sl, Inches(9.1), Inches(1.2), Inches(3.9), Inches(3.5),
         "③ 可插拔 Assigner 模块",
         ["调度与路径之间的核心桥梁",
          "支持基于规则或学习的分配策略",
          "获取路由任务，按规则分配给 AGV 执行",
          "已实现规则性分配算法 + RL 算法"],
         C_BLUE, RGBColor(0xE8, 0xF0, 0xFE))

# 与现有方案的区别
add_rect(sl, Inches(0.5), Inches(5.0), Inches(12.3), Inches(2.0),
         fill_color=RGBColor(0xFD, 0xF2, 0xF2), line_color=C_RED, line_width=Pt(2))
add_textbox(sl, Inches(0.7), Inches(5.1), Inches(12), Inches(0.4),
            "与现有方案的本质区别", font_size=18, bold=True, color=C_RED)
add_bullet_list(sl, Inches(0.7), Inches(5.5), Inches(12), Inches(1.4),
                ["不是「改进 FJSP 或 MAPF 某一端」，而是在环境层面建模双向交互",
                 "Assigner 作为核心耦合接口，是学习的主要目标，填补「搬运衔接无人统筹」的空白",
                 "算法可插拔（任意 FJSP × MAPF 组合一键对比）、耦合可度量（17 项指标）、实验可复现"],
                font_size=14, color=C_DARK, spacing=Pt(6))

# ════════════════════════════════════════════
# Slide 5: 创新点二详解 — 多专家蒸馏
# ════════════════════════════════════════════
sl = add_blank_slide()
add_rect(sl, Inches(0), Inches(0), SW, Inches(0.9), fill_color=C_GREEN)
add_textbox(sl, Inches(0.5), Inches(0.15), Inches(12), Inches(0.6),
            "创新点二：基于 Student-Induced States 的多专家在线蒸馏框架",
            font_size=24, bold=True, color=C_WHITE)

# 问题
add_card(sl, Inches(0.5), Inches(1.2), Inches(6.0), Inches(2.2),
         "现有方法的两重失效",
         ["① 分布偏移：BC 在专家分布上训练，推理时进入新状态区域，误差 O(εT²) 累积放大",
          "② 专家冲突：固定权重平均多专家信号，冲突时产生模糊策略，对所有专家都次优"],
         C_RED, RGBColor(0xFC, 0xE8, 0xE6))

# 理论链
add_rect(sl, Inches(7.0), Inches(1.2), Inches(5.8), Inches(2.2),
         fill_color=RGBColor(0xE6, 0xF4, 0xEA), line_color=C_GREEN, line_width=Pt(2))
add_textbox(sl, Inches(7.2), Inches(1.3), Inches(5.4), Inches(0.4),
            "理论演进链条", font_size=18, bold=True, color=C_GREEN)
add_textbox(sl, Inches(7.2), Inches(1.7), Inches(5.4), Inches(1.5),
            "BC（离线，分布偏移）\n  ↓\nDAgger（在线，单专家，假设 oracle 权威）\n  ↓\nOPD（在线，可质疑 teacher）\n  ↓\n多专家可靠性加权蒸馏（在线，多专家，动态评估）",
            font_size=14, color=C_DARK)

# 核心创新
add_card(sl, Inches(0.5), Inches(3.8), Inches(12.3), Inches(3.2),
         "五大核心创新",
         ["① 蒸馏发生在 student 自己诱导出的状态上（DAgger 式 rollout），从根本上消除分布偏移",
          "② 多专家可靠性动态评估：同一状态下组织多个专家提供候选，通过组内相对比较判断可靠性",
          "③ 状态依赖蒸馏权重：w_k(s) = softmax(agreement_k(s) / τ)，不同状态信任不同专家",
          "④ 在线贝叶斯更新：上线阶段持续估计专家可靠性，动态调整蒸馏权重",
          "⑤ 蒸馏稳定性保障：添加小头部网络，确保微调时蒸馏算法与原算法差异可控"],
         C_GREEN, RGBColor(0xE6, 0xF4, 0xEA))

# ════════════════════════════════════════════
# Slide 6: 创新点三详解 — 多塔策略网络
# ════════════════════════════════════════════
sl = add_blank_slide()
add_rect(sl, Inches(0), Inches(0), SW, Inches(0.9), fill_color=C_PURPLE)
add_textbox(sl, Inches(0.5), Inches(0.15), Inches(12), Inches(0.6),
            "创新点三：专家增强型多塔策略网络", font_size=28, bold=True, color=C_WHITE)

# 问题定位
add_card(sl, Inches(0.5), Inches(1.2), Inches(12.3), Inches(1.5),
         "问题定位",
         ["已有专家模型积累了大量领域知识，但传统调度网络仅依赖原始状态特征，无法充分复用",
          "直接微调面临灾难性遗忘风险",
          "核心问题：如何在 student 自主决策过程中，动态识别并高效复用最相关的专家知识？"],
         C_PURPLE, RGBColor(0xF3, 0xE8, 0xFD))

# 架构：状态塔 vs 专家塔
add_card(sl, Inches(0.5), Inches(3.0), Inches(5.8), Inches(3.0),
         "网络架构",
         ["状态塔：输入 AGV、机器、Job、搬运任务等多维状态特征",
          "专家塔：冻结预训练参数 θ_k，仅接入轻量 Adapter A_k",
          "提取专家内部隐表示 h'_k",
          "",
          "动态路由融合：",
          "h_expert = Σ_k w_k(s) · h'_k",
          "w_k(s) 来自创新点二的蒸馏权重",
          "最终拼接状态特征 + 专家特征 → 输出 Assignment Matrix"],
         C_PURPLE, RGBColor(0xF3, 0xE8, 0xFD))

# 设计要点
add_card(sl, Inches(6.8), Inches(3.0), Inches(6.0), Inches(2.2),
         "设计要点",
         ["专家参数冻结：预训练参数不参与微调，避免灾难性遗忘",
          "轻量 Adapter：仅小规模可训练参数适配，保持稳定性",
          "动态路由融合：权重随状态动态变化，不同场景激活不同专家组合"],
         C_PURPLE, RGBColor(0xF3, 0xE8, 0xFD))

# 与创新点二的关系
add_rect(sl, Inches(6.8), Inches(5.5), Inches(6.0), Inches(1.2),
         fill_color=RGBColor(0xFD, 0xF2, 0xF2), line_color=C_RED, line_width=Pt(2))
add_textbox(sl, Inches(7.0), Inches(5.55), Inches(5.6), Inches(1.1),
            "创新点二回答「当前应相信哪个专家」\n创新点三回答「如何把专家知识融入网络」\nw_k(s) 是两者的桥梁 → 评估-融合闭环",
            font_size=15, bold=False, color=C_RED, align=PP_ALIGN.CENTER)

# ════════════════════════════════════════════
# Slide 7: 总结
# ════════════════════════════════════════════
sl = add_blank_slide()
add_rect(sl, Inches(0), Inches(0), SW, Inches(0.9), fill_color=C_DARK)
add_textbox(sl, Inches(0.5), Inches(0.15), Inches(12), Inches(0.6),
            "总结：三层递进的完整技术体系", font_size=28, bold=True, color=C_WHITE)

# 三个大卡片纵向排列
for i, (title, desc, color, bg) in enumerate([
    ("环境层 — SkyEngine",
     "首个 FJSP-MAPF 闭环统一 RL 环境\n闭环共仿真 | 显式阻塞反馈 | 可插拔 Assigner | 17 项评估指标",
     C_BLUE, RGBColor(0xE8, 0xF0, 0xFE)),
    ("框架层 — 多专家在线蒸馏",
     "Student-Induced 状态蒸馏 | 动态可靠性评估 | 在线贝叶斯更新 | 稳定性保障\n理论链：BC → DAgger → OPD → 多专家可靠性加权蒸馏",
     C_GREEN, RGBColor(0xE6, 0xF4, 0xEA)),
    ("网络层 — 专家增强多塔策略",
     "状态塔 + 专家塔并行 | 冻结参数 + 轻量 Adapter | 动态路由融合\nw_k(s) 桥接框架层与网络层，形成评估-融合闭环",
     C_PURPLE, RGBColor(0xF3, 0xE8, 0xFD)),
]):
    y = Inches(1.2) + Inches(1.9) * i
    add_rect(sl, Inches(0.5), y, Inches(12.3), Inches(1.6),
             fill_color=bg, line_color=color, line_width=Pt(2.5))
    add_textbox(sl, Inches(0.8), y + Inches(0.15), Inches(3.5), Inches(0.4),
                title, font_size=20, bold=True, color=color)
    add_textbox(sl, Inches(0.8), y + Inches(0.55), Inches(11.5), Inches(0.9),
                desc, font_size=15, color=C_DARK)

# 底部一行
add_textbox(sl, Inches(0.5), Inches(7.0), Inches(12.3), Inches(0.4),
            "没有 SkyEngine 就没有闭环信号  |  没有蒸馏框架就无法利用专家知识  |  没有多塔网络就无法高效转化知识",
            font_size=14, color=C_GRAY, align=PP_ALIGN.CENTER)

# ── 保存 ──
out_path = r"C:\Files\Work\Files26\260601天工论文准备\天工系统论文创新点汇报.pptx"
prs.save(out_path)
print(f"PPT saved to: {out_path}")
