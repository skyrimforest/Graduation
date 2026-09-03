"""SkyCausal 阶段汇报 PPT（2026-08-26 定稿口径）— 生成脚本 v2"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

# ── 调色板：BACKGROUND → PRIMARY → ACCENT ──
BG      = RGBColor(0xFF, 0xFF, 0xFF)   # 背景：白
NAVY    = RGBColor(0x1F, 0x3A, 0x5F)   # PRIMARY：深海军蓝
NAVY_D  = RGBColor(0x16, 0x2A, 0x45)   # 封面深底
TINT    = RGBColor(0xEE, 0xF2, 0xF7)   # 主色浅晕
ACCENT  = RGBColor(0xC0, 0x56, 0x21)   # ACCENT：赭橙（克制使用）
TEXT    = RGBColor(0x22, 0x30, 0x3E)
MUTED   = RGBColor(0x5F, 0x63, 0x68)
GREEN   = RGBColor(0x0D, 0x65, 0x2D)
RED     = RGBColor(0xC0, 0x39, 0x2B)
FONT    = "Microsoft YaHei"

W, H = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s


def tb(s, x, y, w, h, lines, size=15, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, spacing=6):
    """lines: str 或 [(text, {size,color,bold}), ...] 每项一段"""
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in tf.paragraphs[0:0]:
        pass
    if isinstance(lines, str):
        lines = [(lines, {})]
    for i, (t, o) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = t
        p.alignment = align
        p.space_after = Pt(spacing)
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    return box


def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, shadow=False):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.adjustments[0] = 0.06
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def chip(s, x, y, w, text, fill=TINT, color=NAVY, size=12.5, bold=True, h=Inches(0.38)):
    c = rect(s, x, y, w, h, fill=fill)
    tf = c.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Emu(0)
    p = tf.paragraphs[0]; p.text = text; p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    return c


def arrow(s, x1, y1, x2, y2, color=NAVY, lw=1.5):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    ln.line.color.rgb = color; ln.line.width = Pt(lw)
    ln.shadow.inherit = False
    return ln


def header(s, kicker, title, dark=False):
    tb(s, Inches(0.6), Inches(0.32), Inches(12.1), Inches(0.3),
       kicker, size=12, color=(ACCENT if not dark else RGBColor(0xE8, 0xA8, 0x7D)), bold=True)
    tb(s, Inches(0.6), Inches(0.62), Inches(12.1), Inches(0.75),
       title, size=27, color=(NAVY if not dark else RGBColor(0xFF, 0xFF, 0xFF)), bold=True)


# ════ S1 封面 ════
s = slide(NAVY_D)
tb(s, Inches(0.9), Inches(1.55), Inches(11.5), Inches(0.4),
   "阶段汇报 · 2026-08-26", size=15, color=RGBColor(0xE8, 0xA8, 0x7D), bold=True)
tb(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.9),
   [("闭环 FJSP–MAPF 中的算子组合切换", {"size": 44, "bold": True, "color": RGBColor(0xFF, 0xFF, 0xFF)}),
    ("限时搜索 · 成本感知决策 · 真值获取与因果分析", {"size": 22, "color": RGBColor(0xC9, 0xD6, 0xE6)})],
   spacing=14)
ln = arrow(s, Inches(0.95), Inches(4.55), Inches(3.4), Inches(4.55), color=ACCENT, lw=2.5)
tb(s, Inches(0.9), Inches(4.85), Inches(11.5), Inches(1.2),
   [("SkyCausal · 论文 v4 主线（已与新口径对齐，16 页编译通过）", {"size": 15, "color": RGBColor(0xC9, 0xD6, 0xE6)}),
    ("依据：总体设计精简版 / 预算受限候选竞赛 / golden trace 预注册 / 切换因果分析三层设计", {"size": 13, "color": RGBColor(0x8F, 0xA3, 0xBC)})],
   spacing=8)

# ════ S2 提纲 ════
s = slide()
header(s, "AGENDA", "汇报提纲")
items = [
    ("01", "问题与证据基线", "闭环中的反事实问题；Phase 7 负结果如何约束设计"),
    ("02", "主线定位与方法", "班子不固定的元级决策；限时组合搜索 × Racing × LCB 切换；真值口径分层"),
    ("03", "因果分析与评估协议", "识别问题与估计量；配对克隆反事实三层设计；双口径记账与 OPE 接口"),
    ("04", "计划 · 实验 · 风险", "T0–T5 里程碑与工作量盘点；E1–E4 实验矩阵；诚实边界"),
]
y = Inches(1.75)
for num, t, d in items:
    tb(s, Inches(0.75), y, Inches(1.1), Inches(0.7), num, size=34, color=ACCENT, bold=True)
    tb(s, Inches(2.0), y + Inches(0.02), Inches(10.6), Inches(0.4), t, size=19, color=NAVY, bold=True)
    tb(s, Inches(2.0), y + Inches(0.42), Inches(10.6), Inches(0.4), d, size=13.5, color=MUTED)
    y += Inches(1.32)

# ════ S3 问题 ════
s = slide()
header(s, "01 · 问题", "切换值不值，只能靠「同状态试一遍」回答")
tb(s, Inches(0.6), Inches(1.55), Inches(5.6), Inches(2.9),
   [("闭环 FJSP–MAPF：调度产生运输需求，拥堵与故障又反过来约束调度。", {"size": 15}),
    ("每个决策点都在问一个反事实问题：", {"size": 15, "bold": True, "color": NAVY}),
    ("此刻换掉正在用的求解器组合，剩余 episode 的终局代价会更低吗？低多少？足以支付切换成本吗？", {"size": 15})],
   spacing=10)
tb(s, Inches(0.6), Inches(4.6), Inches(5.6), Inches(2.2),
   [("离线基准答不了：候选求解器的内部目标 ≠ 闭环终局代价；物理时钟在评估期间持续推进。", {"size": 13.5, "color": MUTED}),
    ("→ 只能冻结状态、在隔离分支上执行候选、比较删失口径下的代价。", {"size": 13.5, "color": ACCENT, "bold": True})],
   spacing=8)
# 右侧分叉图
rect(s, Inches(7.1), Inches(2.9), Inches(1.9), Inches(0.85), fill=NAVY)
tb(s, Inches(7.1), Inches(3.02), Inches(1.9), Inches(0.6), "决策点 s̃", size=15,
   color=RGBColor(0xFF, 0xFF, 0xFF), bold=True, align=PP_ALIGN.CENTER)
rect(s, Inches(9.9), Inches(1.75), Inches(2.9), Inches(0.75), fill=TINT, line=NAVY)
tb(s, Inches(9.9), Inches(1.9), Inches(2.9), Inches(0.5), "KEEP 当前组合", size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
rect(s, Inches(9.9), Inches(4.15), Inches(2.9), Inches(0.75), fill=TINT, line=ACCENT)
tb(s, Inches(9.9), Inches(4.3), Inches(2.9), Inches(0.5), "SWITCH 候选 a", size=14, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
arrow(s, Inches(9.0), Inches(3.15), Inches(9.9), Inches(2.15))
arrow(s, Inches(9.0), Inches(3.5), Inches(9.9), Inches(4.5), color=ACCENT)
arrow(s, Inches(12.8), Inches(2.15), Inches(12.8), Inches(4.4), color=MUTED, lw=1)
tb(s, Inches(9.0), Inches(5.15), Inches(3.9), Inches(1.4),
   [("终局代价 J 对比 → Δ", {"size": 13, "color": MUTED, "bold": True}),
    ("G = Δ − λ·C_switch", {"size": 20, "color": NAVY, "bold": True}),
    ("LCB(G) > τ 才切换", {"size": 14, "color": ACCENT, "bold": True})], spacing=6)

# ════ S4 证据基线 ════
s = slide()
header(s, "01 · 证据基线", "Phase 7：元级周期深控被证伪（936 配对 episodes）")
rows = [
    ("比较", "均值差", "95% CI", "结论"),
    ("周期深搜 vs 一次浅搜", "+16.85", "[11.81, 22.42]", "显著更差"),
    ("周期深搜 vs 静态映射", "+19.25", "[14.57, 24.65]", "显著更差"),
    ("Depth-2 vs Depth-1", "+4.01", "[0.96, 8.03]", "无一改善"),
    ("MCTS vs Uniform 分配", "+4.85", "[1.21, 9.56]", "显著更差"),
    ("温 keep vs 冷重选", "−1050.5", "—", "强支持连续性（72/72 vs 18/72）"),
]
tbl = s.shapes.add_table(len(rows), 4, Inches(0.6), Inches(1.6), Inches(8.6), Inches(3.3)).table
for w, i in zip([Inches(2.9), Inches(1.5), Inches(2.2), Inches(2.0)], range(4)):
    tbl.columns[i].width = w
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = val
        p = cell.text_frame.paragraphs[0]
        p.runs[0].font.name, p.runs[0].font.size = FONT, Pt(13)
        p.runs[0].font.bold = r == 0 or c == 3
        p.runs[0].font.color.rgb = (RGBColor(0xFF, 0xFF, 0xFF) if r == 0
                                    else (RED if "更差" in val else (GREEN if "支持" in val else TEXT)))
        if r == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
tb(s, Inches(9.5), Inches(1.7), Inches(3.3), Inches(3.2),
   [("怎么读这张表", {"size": 15, "bold": True, "color": NAVY}),
    ("深与频繁都错；树分配无效；runtime continuity 是唯一强支持机制。", {"size": 13.5}),
    ("被证伪的是「周期性深度控制」形态，不是元级空间本身——决策点一次浅搜恰是被证据支持的形态。", {"size": 13.5, "color": ACCENT, "bold": True})],
   spacing=10)
tb(s, Inches(0.6), Inches(5.3), Inches(8.6), Inches(0.9),
   [("推论：正确的搜索形态是浅、宽、预算受限的候选竞赛——这就是下一节方法的直接依据。", {"size": 14.5, "color": NAVY, "bold": True})])
tb(s, Inches(0.6), Inches(6.95), Inches(9), Inches(0.3), "数据来源：Phase 7 paper-grade 冻结评估（论文 v4 表 1）", size=11, color=MUTED)

# ════ S5 主线定调 ════
s = slide()
header(s, "02 · 主线定位", "主线定调：班子不固定的元级决策")
rect(s, Inches(0.6), Inches(1.6), Inches(12.1), Inches(1.5), fill=TINT)
tb(s, Inches(0.95), Inches(1.78), Inches(11.4), Inches(1.2),
   [("算子组合不由人定、不靠离线预计算——由决策点上的限时闭环 MCTS 在组合空间在线构造，"
     "经统一 horizon 的竞赛评估，净收益 LCB(G) > τ 才切换，切换显式计价。", {"size": 17, "bold": True, "color": NAVY})])
tb(s, Inches(0.6), Inches(3.4), Inches(12.1), Inches(0.35), "与被证伪形态的四点区别（硬约束，违反即重蹈覆辙）",
   size=14, color=NAVY, bold=True)
cons = [("① 决策点触发", "事件/状态驱动\n不是周期性重搜"),
        ("② anytime 语义", "预算即停·浅层优先\n输出最优+置信状态"),
        ("③ 温接 incumbent", "runtime continuity\nPhase 7 唯一强支持机制"),
        ("④ 竞赛评估", "树只在候选×Tape\n评估结构里，不深树")]
x = Inches(0.6)
for t, d in cons:
    rect(s, x, Inches(3.85), Inches(2.85), Inches(1.5), fill=BG, line=NAVY)
    tb(s, x + Inches(0.15), Inches(4.0), Inches(2.55), Inches(0.35), t, size=14.5, color=NAVY, bold=True)
    tb(s, x + Inches(0.15), Inches(4.4), Inches(2.55), Inches(0.85), d, size=12, color=MUTED)
    x += Inches(3.1)
tb(s, Inches(0.6), Inches(5.75), Inches(12.1), Inches(1.1),
   [("对比基线：", {"size": 14, "bold": True, "color": NAVY}),
    ("① 固定组合（全程不换）  ② 离线预计算静态映射  ③ 一次浅搜（及格线——打不过它就如实报告）", {"size": 14})],
   spacing=6)

# ════ S6 管线 ════
s = slide()
header(s, "02 · 方法", "在线管线：搜索提议与评估，规则拍板")
steps = [("①", "决策状态 s̃", "工厂物理 + 求解器上下文\n+ 近期动态"),
         ("②", "限时组合搜索", "V0：冻结 15 配方池\nV1：组合文法开放构造"),
         ("③", "Racing 评估", "统一 horizon · 同 Tape\n逐轮淘汰 · 温续跑"),
         ("④", "LCB(G) > τ ?", "G = Δ − λ·C_switch\n迟滞/冷却防抖"),
         ("⑤", "闭环收尾", "执行切换或保持\nepisode 配对统计")]
x = Inches(0.55)
for i, (n, t, d) in enumerate(steps):
    c = rect(s, x, Inches(2.0), Inches(2.2), Inches(2.15),
             fill=(NAVY if i in (1, 2) else BG), line=(None if i in (1, 2) else NAVY))
    fg = RGBColor(0xFF, 0xFF, 0xFF) if i in (1, 2) else NAVY
    sub = RGBColor(0xC9, 0xD6, 0xE6) if i in (1, 2) else MUTED
    tb(s, x + Inches(0.12), Inches(2.12), Inches(0.5), Inches(0.4), n, size=17, color=(ACCENT if i not in (1,2) else RGBColor(0xE8,0xA8,0x7D)), bold=True)
    tb(s, x + Inches(0.12), Inches(2.55), Inches(1.96), Inches(0.4), t, size=14.5, color=fg, bold=True)
    tb(s, x + Inches(0.12), Inches(3.0), Inches(1.96), Inches(1.0), d, size=11.5, color=sub)
    if i < 4:
        arrow(s, x + Inches(2.2), Inches(3.05), x + Inches(2.45), Inches(3.05))
    x += Inches(2.45)
rect(s, Inches(0.55), Inches(4.75), Inches(12.2), Inches(1.9), fill=TINT)
tb(s, Inches(0.9), Inches(4.95), Inches(11.5), Inches(1.6),
   [("角色分工（回应「标注器 vs 控制器」）", {"size": 15, "bold": True, "color": NAVY}),
    ("搜索负责提出与评估候选；控制器是成本感知决策规则——搜索不直接拍板。", {"size": 14}),
    ("已登记缺口：Racing 当前仅排序、未显式计价 C_switch，T1 成本表接入后 G 才完整（T3 前补上）。", {"size": 13, "color": ACCENT})],
   spacing=8)

# ════ S7 Racing ════
s = slide()
header(s, "02 · 方法", "Racing：候选竞赛评估器（已实现 · 102 项测试套件内）")
tb(s, Inches(0.6), Inches(1.6), Inches(5.7), Inches(3.4),
   [("机制要点", {"size": 15, "bold": True, "color": NAVY}),
    ("统一 horizon 铁律：同一轮所有候选跑一样远、用同一批 Tape——预算绝不制造 per-candidate 偏差", {"size": 13.5}),
    ("温续跑：旧分支从上次子状态增量推进 O(ΔH)，保持 runtime continuity", {"size": 13.5}),
    ("逐轮淘汰：总成本几何级数收敛于首轮全宽成本，省一个量级", {"size": 13.5}),
    ("三态诚实记录：终局 / 删失 / 求解器错误分开统计，删失罚值不冒充真值", {"size": 13.5})],
   spacing=9)
rect(s, Inches(6.7), Inches(1.6), Inches(6.0), Inches(3.4), fill=TINT)
tb(s, Inches(7.0), Inches(1.78), Inches(5.4), Inches(3.1),
   [("穷举退化（命题）", {"size": 15, "bold": True, "color": NAVY}),
    ("不淘汰 + 全 Tape + 跑到终局 ⇒ 退化为完全枚举，输出冻结候选×Tape 集上的精确 Q̂", {"size": 13.5}),
    ("anytime 性质：淘汰只删统计上被支配的候选，保留最优的概率不减——预算耗尽时输出当前置信与删失状态，不伪造精度", {"size": 13.5, "color": ACCENT, "bold": True})],
   spacing=10)
tb(s, Inches(0.6), Inches(5.3), Inches(12.1), Inches(0.35), "实跑示例（tiny 闭环场面，Docker 编排，宿主机/容器逐字节一致）",
   size=14, color=NAVY, bold=True)
demo = [("轮次", "H=20", "H=60", "H=120"), ("存活", "8→4", "4→3", "3→2")]
dt = s.shapes.add_table(2, 4, Inches(0.6), Inches(5.7), Inches(6.4), Inches(0.85)).table
for r in range(2):
    for c in range(4):
        cell = dt.cell(r, c)
        cell.text = demo[r][c]
        p = cell.text_frame.paragraphs[0]
        p.runs[0].font.name, p.runs[0].font.size = FONT, Pt(12.5)
        p.runs[0].font.bold = r == 0
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r == 0 else TEXT
        if r == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
tb(s, Inches(7.4), Inches(5.7), Inches(5.3), Inches(0.9),
   [("keep 以 76.5 被 adaptive 的 73.5 淘汰；top-3 = 72 / 72 / 73；全程 50 次转移、约 11 秒。", {"size": 12.5, "color": MUTED})])

# ════ S8 真值口径 ════
s = slide()
header(s, "02 · 真值", "真值口径分层：一台主力 + 一台验证器")
cards = [
    ("主力 · golden trace", "每 spec K=8 条随机闭环轨迹，取终局最优作参考值（72 spec）", "语义边界：K 内最优 ≠ 全局最优；称「参考值」，不冒充穷举真值或公平基线", NAVY, RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xC9, 0xD6, 0xE6)),
    ("验证器 · Racing 穷举", "冻结候选 × 冻结 Tape 集上的精确 Q̂（穷举退化命题）", "服务 fixed 组合基线确认 + 小子集上限时排序的一致性校验", BG, NAVY, MUTED),
    ("future · oracle 标签", "逐决策点 Δ/G 反事实重算 + 预测器训练", "已移出本文范围——在线值由 Racing 估计 + 成本表计价", BG, MUTED, MUTED),
]
y = Inches(1.65)
for t, d1, d2, fill, fg, sub in cards:
    rect(s, Inches(0.6), y, Inches(12.1), Inches(1.55), fill=fill, line=(None if fill == NAVY else MUTED))
    tb(s, Inches(0.95), y + Inches(0.14), Inches(3.2), Inches(0.4), t, size=16, color=fg, bold=True)
    tb(s, Inches(4.3), y + Inches(0.14), Inches(8.1), Inches(0.55), d1, size=13.5, color=fg, bold=True)
    tb(s, Inches(4.3), y + Inches(0.72), Inches(8.1), Inches(0.7), d2, size=12, color=sub)
    y += Inches(1.78)
tb(s, Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.35),
   "三种口径并存——任何场合混用即构成过度声明（风险清单第 2 条）", size=13.5, color=ACCENT, bold=True)

# ════ S9 因果-识别 ════
s = slide()
header(s, "03 · 因果分析", "因果分析 I：先问「效应」到底是什么")
tb(s, Inches(0.6), Inches(1.6), Inches(5.8), Inches(3.2),
   [("观测性前后对比的三重隐患", {"size": 15, "bold": True, "color": NAVY}),
    ("选择偏差：自适应系统只在自认为有利时才换", {"size": 14}),
    ("时间趋势：指标本来就在变", {"size": 14}),
    ("均值回归：往往指标很差才触发切换，之后自然回弹", {"size": 14}),
    ("→ 自适应日志的事件研究只能描述，不能下因果结论。", {"size": 13.5, "color": ACCENT, "bold": True})],
   spacing=9)
rect(s, Inches(6.8), Inches(1.6), Inches(5.9), Inches(3.2), fill=TINT)
tb(s, Inches(7.1), Inches(1.78), Inches(5.3), Inches(2.9),
   [("估计量先行", {"size": 15, "bold": True, "color": NAVY}),
    ("切换的效应依赖于切换之后的后续策略 π⁺——这是随机化也消不掉的定义性问题。", {"size": 13.5}),
    ("τ_π⁺(s̃,a) = J_π⁺(keep) − J_π⁺(a)", {"size": 17, "bold": True, "color": ACCENT}),
    ("主口径：switch-then-keep（机制最干净）；随机继续 = 与自然实验可比的桥。", {"size": 13})],
   spacing=9)
tb(s, Inches(0.6), Inches(5.15), Inches(12.1), Inches(1.5),
   [("我们的优势：有可克隆、可重放、确定性的仿真器。", {"size": 16, "bold": True, "color": NAVY}),
    ("无仿真器世界只能在识别假设下「估计」反事实（合成控制/事件研究类）；我们让反事实被仿真「实现」——识别假设全部短路。", {"size": 14})],
   spacing=8)

# ════ S10 因果三层 ════
s = slide()
header(s, "03 · 因果分析", "因果分析 II：三层识别设计")
layers = [
    ("L2 主力", "配对克隆反事实", "决策根克隆 + 同一 Tape 跑 keep/switch 至终局 → 个体级配对差分。"
     "同 Tape = 公共随机数方差缩减；反事实被实现而非估计，不依赖不可检验的识别假设。"
     "分层 → CATE（何时换）；通道归因 + 消融 → 为何换。", NAVY, RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xD5, 0xDF, 0xEA)),
    ("L1 交叉验证", "随机轨迹自然实验", "round-1 切换均匀随机，天然免除选择偏差；无偏但功效有限。"
     "与 L2 识别原理不同——结论方向一致即稳健性证据。", TINT, NAVY, TEXT),
    ("L0 描述", "自适应日志事件研究", "切换事件前后窗口的聚合指标轨迹；显式标注选择偏差，只作描述。", BG, MUTED, MUTED),
]
y = Inches(1.6)
for tag, t, d, fill, fg, sub in layers:
    rect(s, Inches(0.6), y, Inches(12.1), Inches(1.45), fill=fill, line=(None if fill == NAVY else (NAVY if fill == TINT else MUTED)))
    tb(s, Inches(0.9), y + Inches(0.12), Inches(1.7), Inches(0.4), tag, size=14, color=(ACCENT if fill == NAVY else fg), bold=True)
    tb(s, Inches(0.9), y + Inches(0.52), Inches(2.6), Inches(0.6), t, size=16.5, color=fg, bold=True)
    tb(s, Inches(3.8), y + Inches(0.15), Inches(8.6), Inches(1.15), d, size=12.5, color=sub)
    y += Inches(1.62)
tb(s, Inches(0.6), Inches(6.55), Inches(3.0), Inches(0.4), "诚实边界：", size=13.5, color=NAVY, bold=True)
for i, c in enumerate(["保真度前置门为前提", "CP-SAT 非确定 → 排除或增 Tape", "π⁺ 显式定义"]):
    chip(s, Inches(2.4) + Inches(i * 3.55), Inches(6.55), Inches(3.4), c, size=12)

# ════ S11 双口径 + OPE ════
s = slide()
header(s, "03 · 评估协议", "双口径记账 + 给 OPE 留门")
rect(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(4.6), fill=TINT)
tb(s, Inches(0.95), Inches(1.8), Inches(5.2), Inches(4.2),
   [("双口径：质量与开销分开，不合成", {"size": 16, "bold": True, "color": NAVY}),
    ("质量轴：makespan（整数仿真步）", {"size": 14.5}),
    ("开销轴：决策次数 · 墙钟 · 转移调用 · 仿真步数", {"size": 14.5}),
    ("所有方法两轴分别配对比较；λ 只是决策规则内部交换率，敏感性单独报告，不造加权合成分数。", {"size": 13.5}),
    ("公平性声明：固定基线几乎不付搜索开销，我们的决策期计算成本全部计入开销轴。", {"size": 13, "color": ACCENT})],
   spacing=11)
rect(s, Inches(6.8), Inches(1.6), Inches(5.9), Inches(4.6), fill=BG, line=NAVY)
tb(s, Inches(7.15), Inches(1.8), Inches(5.2), Inches(4.2),
   [("OPE 接口：现在埋，未来才可用", {"size": 16, "bold": True, "color": NAVY}),
    ("未来（含 LLM 调度）想从策略自己的日志估计「换一个策略会怎样」→ 离策略评估（OPE / doubly-robust）。", {"size": 13.5}),
    ("死穴：日志中每个动作的执行概率必须已知且非零——确定性策略的日志对 OPE 不可用。", {"size": 13.5, "color": RED, "bold": True}),
    ("动作：shadow 决策注入 ε 随机化 + 记录动作概率。成本近零；不做，未来只能重跑全量实验。", {"size": 13.5, "color": ACCENT, "bold": True})],
   spacing=11)
tb(s, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.5),
   "策略级总体效应不走观测层——主实验（72 spec × 同 Tape 配对）本身就是随机对照实验。", size=14, color=NAVY, bold=True)

# ════ S12 里程碑 + 工作量 ════
s = slide()
header(s, "04 · 计划", "里程碑 T0–T5 与工作量盘点")
arrow(s, Inches(0.8), Inches(2.35), Inches(12.6), Inches(2.35), color=MUTED, lw=1.5)
nodes = [("T0 保真门", "半天"), ("T1 组合契约", "1 天"), ("T2 限时搜索", "V0→V1 · 2–3 天"),
         ("T3 LCB 切换", "1 天 · shadow"), ("T5 闭环评估", "实验矩阵")]
x = Inches(0.75)
for i, (t, d) in enumerate(nodes):
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.55), Inches(2.22), Inches(0.26), Inches(0.26))
    dot.fill.solid(); dot.fill.fore_color.rgb = ACCENT; dot.line.fill.background(); dot.shadow.inherit = False
    tb(s, x, Inches(1.62), Inches(2.3), Inches(0.4), t, size=14.5, color=NAVY, bold=True)
    tb(s, x, Inches(2.62), Inches(2.3), Inches(0.4), d, size=12, color=MUTED)
    x += Inches(2.4)
rect(s, Inches(5.4), Inches(3.25), Inches(2.6), Inches(0.75), fill=BG, line=GREEN)
tb(s, Inches(5.5), Inches(3.33), Inches(2.4), Inches(0.6),
   [("T4 golden trace round-1", {"size": 13, "bold": True, "color": GREEN}),
    ("过 T0 即并行开跑，不等 hero", {"size": 11.5, "color": MUTED})], spacing=2)
tb(s, Inches(0.6), Inches(4.35), Inches(12.1), Inches(0.35), "工作量真相：复杂的是存量系统，不是剩余增量", size=15, color=NAVY, bold=True)
facts = [("已建成", "Racing / mcts_node / uniform / trigger_policy / admission / 闭环控制器——102 项测试兜底", GREEN),
         ("唯一新代码", "V1 组合文法候选生成器（配方空间本就按维度结构化，合法性掩码已就位）", ACCENT),
         ("总体增量", "集中做约 2 周 + 机器时间；V0 ≈ 已有 Racing + LCB，可最先跑通、先锁论文下限", NAVY)]
y = Inches(4.8)
for t, d, c in facts:
    tb(s, Inches(0.6), y, Inches(1.9), Inches(0.5), t, size=13.5, color=c, bold=True)
    tb(s, Inches(2.6), y, Inches(10.1), Inches(0.55), d, size=13)
    y += Inches(0.62)

# ════ S13 实验矩阵 ════
s = slide()
header(s, "04 · 实验", "实验矩阵：8 张表 · 7 张图")
exps = [
    ("E1 主闭环对比", "72 spec × {固定, 静态映射, 一次浅搜, V0, V1} 严格配对", "主结果表 / 帕累托图 / V0-V1 分离"),
    ("E2 golden round-1", "72 spec × K=8 随机轨迹（协议已冻结）", "参考值表 / 因果自然实验面板"),
    ("E2a 配对克隆反事实", "决策根 × {keep, switch} × 同 Tape → 终局", "CATE 表 / 通道归因图（何时换·为何换）"),
    ("E3 敏感性 / E4 校验", "λ·τ·预算网格；小子集穷举 vs 限时排序", "敏感性曲线 / 一致性表"),
]
y = Inches(1.7)
for t, d, out in exps:
    rect(s, Inches(0.6), y, Inches(12.1), Inches(1.12), fill=(TINT if t.startswith("E2a") else BG),
         line=(None if t.startswith("E2a") else MUTED))
    tb(s, Inches(0.9), y + Inches(0.08), Inches(3.3), Inches(0.9), t, size=15, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, Inches(4.3), y + Inches(0.08), Inches(4.9), Inches(0.9), d, size=13, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, Inches(9.3), y + Inches(0.08), Inches(3.2), Inches(0.9), out, size=12.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(1.28)
tb(s, Inches(0.6), Inches(6.95), Inches(12.1), Inches(0.35),
   "新图脚本沿用 manifest+哈希纪律（generate_v5_figures.py）", size=12, color=MUTED)

# ════ S14 风险 ════
s = slide()
header(s, "04 · 风险", "风险与诚实边界")
risks = [
    ("① 深控重演", "元级搜索重蹈 Phase 7——防线 = 四条硬约束，及格线 = 一次浅搜，不合格即如实报告"),
    ("② 真值语义漂移", "best-of-K 参考值 / 冻结集穷举值 / 在线估计三种口径并存，任何场合混用即过度声明"),
    ("③ 因果分析边界", "配对层上限 = 仿真保真度（CP-SAT 非确定需排除）；自然实验只覆盖随机策略可达分布；事件研究带选择偏差——三层范围不得互换"),
    ("④ V0/V1 口径分离", "开放组合空间的收益必须与「自适应本身」的收益分离报告，禁止混表"),
    ("⑤ baseline 公平性", "LLM 对比阶段同环境、同动作空间、同预算，禁止给任何一方更宽松的条件"),
]
y = Inches(1.6)
for t, d in risks:
    tb(s, Inches(0.6), y, Inches(2.6), Inches(0.9), t, size=15, color=NAVY, bold=True)
    tb(s, Inches(3.4), y, Inches(9.3), Inches(0.95), d, size=13.5)
    arrow(s, Inches(0.6), y + Inches(0.98), Inches(12.7), y + Inches(0.98), color=RGBColor(0xE2, 0xE8, 0xF0), lw=1)
    y += Inches(1.06)

# ════ S15 总结 ════
s = slide(NAVY_D)
tb(s, Inches(0.9), Inches(0.7), Inches(11.5), Inches(0.4), "总结", size=15,
   color=RGBColor(0xE8, 0xA8, 0x7D), bold=True)
tb(s, Inches(0.9), Inches(1.15), Inches(11.5), Inches(1.5),
   [("组合由搜索产生 · 切换显式计价 · 真值分层声明 · 质量与开销分开记账", {"size": 27, "bold": True, "color": RGBColor(0xFF, 0xFF, 0xFF)})])
cols = [
    ("已交付", ["反事实仿真基础设施（已独立确认）", "元级周期深控证伪（936 episodes）",
                "Racing 评估器 + 真值性质（已实现）", "论文 v4 对齐新主线（16 页编译通过）"]),
    ("下一步", ["T0 保真门（半天）→ T4 golden trace 开跑", "T1 契约表 → T2 V0 先行（先锁下限）",
                "E2a 配对克隆反事实（因果章主力）", "shadow 加 ε 随机化并记录概率"]),
]
x = Inches(0.9)
for t, items in cols:
    tb(s, x, Inches(3.1), Inches(5.6), Inches(0.4), t, size=17,
       color=RGBColor(0xE8, 0xA8, 0x7D), bold=True)
    lines = [(("· " + it), {"size": 14, "color": RGBColor(0xC9, 0xD6, 0xE6)}) for it in items]
    tb(s, x, Inches(3.6), Inches(5.6), Inches(2.4), lines, spacing=8)
    x += Inches(6.0)
tb(s, Inches(0.9), Inches(6.6), Inches(11.5), Inches(0.4),
   "SkyCausal · 阶段汇报 · 2026-08-26", size=12, color=RGBColor(0x8F, 0xA3, 0xBC))

OUT = "SkyCausal_阶段汇报_20260826.pptx"
prs.save(OUT)
print("saved", OUT)
print("slides:", len(prs.slides._sldIdLst))
